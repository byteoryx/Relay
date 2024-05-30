from openpyxl.styles import Color, PatternFill, Font, Alignment, Border, Side
from openpyxl import Workbook, load_workbook
from datetime import datetime
from time import sleep
import os

from .utils import logger
from .wallet import Wallet


class Excel:
    def __init__(self, total_len: int, name: str):
        if not os.path.isdir('results'): os.mkdir('results')
        self.name = name

        workbook = Workbook()
        sheet = workbook.active
        self.file_name = f'{name}_{total_len}accs_{datetime.now().strftime("%d_%m_%Y_%H_%M_%S")}.xlsx'

        sheet['A1'] = 'EVM privatekey'
        sheet['B1'] = 'EVM Address'
        sheet['C1'] = 'Status'

        sheet.column_dimensions['A'].width = 15
        sheet.column_dimensions['B'].width = 46
        sheet.column_dimensions['C'].width = 25

        for cell in sheet._cells:
            sheet.cell(cell[0], cell[1]).font = Font(bold=True)
            sheet.cell(cell[0], cell[1]).alignment = Alignment(horizontal='center')
            sheet.cell(cell[0], cell[1]).border = Border(left=Side(style='thin'), bottom=Side(style='thin'), right=Side(style='thin'))

        workbook.save('results/'+self.file_name)


    def edit_table(self, wallet: Wallet, status: str):
        while True:
            try:
                workbook = load_workbook('results/'+self.file_name)
                sheet = workbook.active

                valid_info = [
                    wallet.privatekey,
                    wallet.address,
                    str(status),
                ]
                sheet.append(valid_info)

                for row_cells in sheet.iter_rows(min_row=sheet.max_row, max_row=sheet.max_row):
                    for cell in row_cells:
                        cell.border = Border(left=Side(style='thin'), right=Side(style='thin'))
                        if cell.column == 3:
                            if str(cell.value) == "True": rgb_color = '32CD32'
                            else: rgb_color = 'ff0f0f'
                            cell.fill = PatternFill(patternType='solid', fgColor=Color(rgb=rgb_color))

                workbook.save('results/'+self.file_name)
                return True
            except PermissionError:
                logger.warning(f'Excel | Cant save excel file, close it!')
                sleep(3)
            except Exception as err:
                logger.critical(f'Excel | Cant save excel file: {err} | {wallet.address}')
                return False
